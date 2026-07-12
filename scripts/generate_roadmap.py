import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def create_roadmap():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Yol Haritası"

    # Grid çizgilerini görünür kıl
    ws.views.sheetView[0].showGridLines = True

    # Başlık
    ws["A1"] = "Multi-CMS Manager Geliştirici Yol Haritası"
    ws["A1"].font = Font(name="Segoe UI", size=16, bold=True, color="2C3E50")
    ws.merge_cells("A1:E1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    # Başlıklar
    headers = ["Aşama", "Süre", "Durum", "İlerleme", "Detaylı Görevler"]
    ws.append([])  # Boş satır
    ws.append(headers)

    # Başlık Hücre Tasarımı
    header_fill = PatternFill(
        start_color="34495E", end_color="34495E", fill_type="solid",
    )
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(
        horizontal="center", vertical="center", wrap_text=True,
    )

    for col_idx in range(1, 6):
        cell = ws.cell(row=3, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # Yol Haritası Verisi
    stages = [
        (
            "AŞAMA 0: Altyapı ve Mimari Temel",
            "Hafta 1-2",
            "Tamamlandı",
            1.00,
            "Klasör hiyerarşisi, pyproject.toml bağımlılıkları, Pydantic Settings config, SQLAlchemy & SQLite/SQLCipher bağlantı katmanı, structlog loglama, pytest & Mock API sunucu altyapısı, ruff ve GitHub Actions CI kurulumu.",
        ),
        (
            "AŞAMA 1: Yedekleme Motoru (Core)",
            "Hafta 3-5",
            "Tamamlandı",
            1.00,
            "BaseCMSAdapter tasarımı, Dolibarr & WordPress/WooCommerce adaptörleri, DatabaseBackupService (mysqldump), FileBackupService (tar.gz), BackupOrchestrator, SFTP transferi, Görev Zamanlayıcı entegrasyonu.",
        ),
        (
            "AŞAMA 2: Yerel Veritabanı ve API Bağdaştırıcıları",
            "Hafta 6-8",
            "Tamamlandı",
            1.00,
            "Site, Product, Order, SyncLog, ChangeLog tabloları. Soft Delete, Optimistic Locking, SQL Listeners ile ChangeLog (PENDING_PUSH). Dolibarr & WooCommerce client'lar, API key için keyring entegrasyonu.",
        ),
        (
            "AŞAMA 3: Tek Yönlü Senkronizasyon — Pull",
            "Hafta 9-10",
            "Devam Ediyor",
            0.50,
            "SyncEngine Pull yapısı, Delta Senkronizasyon (modified_after), sayfalama (Pagination Iterator), akıllı veri birleştirme (Merge) stratejisi, sync_logs audit kaydı, CLI sync pull testleri.",
        ),
        (
            "AŞAMA 4: Çift Yönlü Senkronizasyon — Push & Çatışma Yönetimi",
            "Hafta 11-14",
            "Planlandı",
            0.00,
            "PushEngine PENDING_PUSH uzak API gönderimi, Conflict Detection (remote_modified_at > local_last_sync), Conflict Resolution (Last Write Wins vb.), tenacity backoff, pybreaker circuit breaker, PriorityQueue SyncQueue.",
        ),
        (
            "AŞAMA 5: Masaüstü Arayüzü — PyQt6",
            "Hafta 15-18",
            "Planlandı",
            0.00,
            "QStackedWidget ana pencere, Site Yönetimi şifreli girişler, QTableView lazy loading Ürün listesi, Sipariş yönetimi durum güncellemeleri, Canlı log ve çatışma Modal seçim ekranı, asenkron QThreadPool thread yönetimi.",
        ),
        (
            "AŞAMA 6: Gelişmiş Özellikler",
            "Hafta 19-21",
            "Planlandı",
            0.00,
            "pandas/openpyxl ile Excel/PDF raporları, Dolibarr Multicompany/WooCommerce Multisite, Düşük stok Toast bildirimleri, Ctrl+K hızlı arama, kısayollar, Excel toplu içe aktarım (data_io), QTimer 15dk otomatik sync.",
        ),
        (
            "AŞAMA 7: Paketleme ve Dağıtım",
            "Hafta 22-23",
            "Planlandı",
            0.00,
            "PyInstaller build.py, Inno Setup (.exe) / macOS .app / Linux AppImage, GitHub Releases API Auto-updater, İlk Kurulum Sihirbazı (db key oluşturma), teknik ve son kullanıcı dokümantasyonları.",
        ),
    ]

    # Hücre kenarlıkları
    thin_border = Border(
        left=Side(style="thin", color="BDC3C7"),
        right=Side(style="thin", color="BDC3C7"),
        top=Side(style="thin", color="BDC3C7"),
        bottom=Side(style="thin", color="BDC3C7"),
    )

    # Verileri ekle ve şekillendir
    row_idx = 4
    for stage, duration, status, progress, details in stages:
        ws.append([stage, duration, status, progress, details])

        # Satır arkaplan rengi (zebra striping)
        row_fill = PatternFill(
            start_color="F8F9FA" if row_idx % 2 == 0 else "FFFFFF",
            end_color="F8F9FA" if row_idx % 2 == 0 else "FFFFFF",
            fill_type="solid",
        )

        # Durum sütunu renklendirmeleri
        status_fill = None
        status_font_color = "000000"
        if status == "Tamamlandı":
            status_fill = PatternFill(
                start_color="D4EFDF", end_color="D4EFDF", fill_type="solid",
            )
            status_font_color = "196F3D"
        elif status == "Devam Ediyor":
            status_fill = PatternFill(
                start_color="FCF3CF", end_color="FCF3CF", fill_type="solid",
            )
            status_font_color = "B7950B"
        else:
            status_fill = PatternFill(
                start_color="E5E7E9", end_color="E5E7E9", fill_type="solid",
            )
            status_font_color = "5D6D7E"

        for col_idx in range(1, 6):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            cell.fill = row_fill

            # Hizalama ve biçimlendirmeler
            if col_idx == 1:
                cell.font = Font(
                    name="Segoe UI", size=10, bold=True, color="2C3E50",
                )
            elif col_idx in (2, 3, 4):
                cell.alignment = Alignment(
                    horizontal="center", vertical="center",
                )
            elif col_idx == 5:
                cell.alignment = Alignment(wrap_text=True, vertical="center")

            if col_idx == 3:  # Durum sütunu stil ezmesi
                cell.fill = status_fill
                cell.font = Font(
                    name="Segoe UI", size=10, bold=True, color=status_font_color,
                )
            elif col_idx == 4:  # Yüzdelik format
                cell.number_format = "0%"

        row_idx += 1

    # Sütun genişliklerini ayarla
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col_letter == "A":
            ws.column_dimensions[col_letter].width = 45
        elif col_letter == "B":
            ws.column_dimensions[col_letter].width = 15
        elif col_letter == "C":
            ws.column_dimensions[col_letter].width = 18
        elif col_letter == "D":
            ws.column_dimensions[col_letter].width = 12
        elif col_letter == "E":
            ws.column_dimensions[col_letter].width = 90

    # Satır yüksekliklerini ayarla
    ws.row_dimensions[3].height = 25
    ws.row_dimensions[1].height = 35

    wb.save("Multi_CMS_Manager_Gelistirici_Yol_Haritasi.xlsx")
    print("Multi_CMS_Manager_Gelistirici_Yol_Haritasi.xlsx başarıyla üretildi!")


if __name__ == "__main__":
    create_roadmap()
