import difflib
import os
from typing import Any

import openpyxl
import pandas as pd
from sqlalchemy.orm import Session

from src.core.models import ChangeLog, Customer, Product

# ==========================================
# ALAN TANIMLARI VE ÇEVİRİ SÖZLÜKLERİ
# ==========================================

CUSTOMER_FIELDS = {
    "customer_code": {"display": "Cari Kodu", "aliases": ["cari kod", "kod", "müşteri no", "code_client", "client_code"]},
    "fullname": {"display": "Ticari Unvanı", "aliases": ["ünvan", "ad soyad", "müşteri adı", "ticari unvan", "nom", "name", "fullname"]},
    "email": {"display": "E-Posta", "aliases": ["e-posta", "email", "posta", "e-mail", "mail"]},
    "phone": {"display": "Telefon", "aliases": ["telefon", "tel", "phone", "gsm", "mobil"]},
    "address": {"display": "Adres", "aliases": ["adres", "address", "lokasyon"]},
    "tax_office": {"display": "Vergi Dairesi", "aliases": ["vergi dairesi", "v.d.", "dairesi", "tax_office", "localtax1_ass"]},
    "tax_number": {"display": "Vergi Numarası", "aliases": ["vergi no", "tckn", "vkn", "vergi numarası", "tax_number", "tva_intra", "tva_ass"]},
    "group_name": {"display": "Grubu", "aliases": ["grubu", "grup", "group", "group_name"]},
    "sub_group_1": {"display": "Ara Grubu", "aliases": ["ara grubu", "ara grup", "sub_group_1"]},
    "sub_group_2": {"display": "Alt Grubu", "aliases": ["alt grubu", "alt grup", "sub_group_2"]},
    "special_code_1": {"display": "Özel Kodu 1", "aliases": ["özel kodu 1", "özel kod 1", "special_code_1"]},
    "special_code_2": {"display": "Özel Kodu 2", "aliases": ["özel kodu 2", "özel kod 2", "special_code_2"]},
    "special_code_3": {"display": "Özel Kodu 3", "aliases": ["özel kodu 3", "özel kod 3", "special_code_3"]},
}

PRODUCT_FIELDS = {
    "sku": {"display": "Ürün Kodu", "aliases": ["ürün kodu", "ref", "sku", "kod", "product_code"]},
    "name": {"display": "Ürün Adı", "aliases": ["ürün adı", "label", "name", "ad", "product_name"]},
    "description": {"display": "Açıklama", "aliases": ["açıklama", "description", "detay", "info"]},
    "base_price": {"display": "Satış Fiyatı (KDV Hariç)", "aliases": ["fiyat kdv hariç", "base price", "base_price", "fiyat_kdv_haric"]},
    "price": {"display": "Satış Fiyatı", "aliases": ["satış fiyatı", "fiyat", "price", "satış_fiyatı"]},
    "stock": {"display": "Stok Miktarı", "aliases": ["stok", "stock", "miktar", "qty", "stock_real", "quantity"]},
    "barcode": {"display": "Barkod", "aliases": ["barkod", "barcode"]},
    "vat_rate": {"display": "KDV Oranı (%)", "aliases": ["kdv", "kdv oranı", "vat_rate", "vat", "kdv_orani"]},
    "status": {"display": "Satış Durumu", "aliases": ["satış durumu", "status", "aktif", "active"]},
    "status_buy": {"display": "Alış Durumu", "aliases": ["alış durumu", "status_buy"]},
}


# ==========================================
# SÜTUN VE HARF DÖNÜŞÜM YARDIMCILARI
# ==========================================

def col_idx_to_letter(idx: int) -> str:
    """Sütun indeks numarasını (0, 1, 2...) Excel sütun harfine (A, B, C...) dönüştürür."""
    letter = ""
    while idx >= 0:
        letter = chr(idx % 26 + 65) + letter
        idx = idx // 26 - 1
    return letter


def letter_to_col_idx(letter: str) -> int:
    """Excel sütun harfini (A, B, C...) sütun indeks numarasına (0, 1, 2...) dönüştürür."""
    idx = 0
    for char in str(letter).upper().strip():
        idx = idx * 26 + (ord(char) - 64)
    return idx - 1


# ==========================================
# EŞLEŞTİRME VE TEMPLATE YARDIMCILARI
# ==========================================

def find_best_match(excel_header: str, fields_dict: dict[str, dict[str, Any]]) -> str | None:
    """Excel başlığı ile veritabanı alan tanımları arasındaki en iyi eşleşmeyi bulur (Fuzzy Matching)."""
    excel_header_clean = str(excel_header).strip().lower()
    if not excel_header_clean:
        return None
    
    # 1. Birebir Display Name veya Aliases tam eşleşme kontrolü
    for field_name, info in fields_dict.items():
        if excel_header_clean == info["display"].lower():
            return field_name
        for alias in info["aliases"]:
            if excel_header_clean == alias.lower():
                return field_name
                
    # 2. Fuzzy benzerlik eşleşmesi
    best_field = None
    best_score = 0.0
    
    for field_name, info in fields_dict.items():
        # Display name benzerliği
        score = difflib.SequenceMatcher(None, excel_header_clean, info["display"].lower()).ratio()
        if score > best_score:
            best_score = score
            best_field = field_name
            
        # Aliases benzerliği
        for alias in info["aliases"]:
            score = difflib.SequenceMatcher(None, excel_header_clean, alias.lower()).ratio()
            if score > best_score:
                best_score = score
                best_field = field_name
                
    # Eşleşme skoru %70 ve üzeri ise kabul et
    if best_score >= 0.70:
        return best_field
        
    return None


def generate_template_file(file_path: str, fields_dict: dict[str, dict[str, Any]]) -> None:
    """Kolon başlıkları Türkçe display name'lerden oluşan boş bir Excel şablonu üretir."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aktarım Şablonu"
    
    headers = [info["display"] for info in fields_dict.values()]
    ws.append(headers)
    
    # Header formatlama (Görsel bütünlük için kalın ve gri dolgulu)
    header_font = openpyxl.styles.Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = openpyxl.styles.PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        
    wb.save(file_path)


def export_to_excel_file(file_path: str, fields_dict: dict[str, dict[str, Any]], data_list: list[Any]) -> None:
    """Verilen ORM nesne listesini Türkçe kolon başlıkları ile Excel formatında dışarı aktarır."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dışa Aktarılan Veriler"
    
    headers = [info["display"] for info in fields_dict.values()]
    ws.append(headers)
    
    # Header stil
    header_font = openpyxl.styles.Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = openpyxl.styles.PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        
    # Satırlar
    for obj in data_list:
        row = []
        for field_name in fields_dict.keys():
            val = getattr(obj, field_name, "")
            row.append(val if val is not None else "")
        ws.append(row)
        
    wb.save(file_path)


# ==========================================
# İÇE AKTARIM MOTORU (IMPORT ENGINE)
# ==========================================

def read_excel_or_ods(file_path: str, sheet_index: int = 1, has_headers: bool = True) -> tuple[list[dict[str, Any]], pd.DataFrame]:
    """Excel veya ODS dosyasını okur. 
    Geriye sütun listesini (harf, başlık, display bilgisiyle) ve DataFrame'i döndürür.
    """
    _, ext = os.path.splitext(file_path.lower())
    pandas_sheet = max(0, sheet_index - 1)
    
    # pandas'a header vermeden ham olarak okuyoruz, başlık satırını kendimiz yöneteceğiz
    if ext == ".ods":
        df = pd.read_excel(file_path, sheet_name=pandas_sheet, engine="odf", header=None)
    else:
        df = pd.read_excel(file_path, sheet_name=pandas_sheet, engine="openpyxl", header=None)
        
    columns_info = []
    headers_row = {}
    start_row = 0
    
    # Eğer ilk satır başlık içeriyorsa
    if has_headers and len(df) > 0:
        for col_idx in range(df.shape[1]):
            val = df.iloc[0, col_idx]
            val_str = str(val).strip() if pd.notna(val) else ""
            headers_row[col_idx] = val_str
        start_row = 1
    else:
        for col_idx in range(df.shape[1]):
            headers_row[col_idx] = ""
            
    # DataFrame'i gerçek veri satırlarıyla güncelle ve sütun indekslerini 0-based yap
    df_data = df.iloc[start_row:].copy()
    df_data.columns = [col_idx for col_idx in range(df.shape[1])]
    
    for col_idx in range(df.shape[1]):
        letter = col_idx_to_letter(col_idx)
        header_text = headers_row[col_idx]
        
        # Eğer pandas otomatik isimlendirdiyse veya boşsa
        if header_text and not str(header_text).startswith("Unnamed:"):
            display_text = f"{letter} - {header_text}"
        else:
            display_text = f"{letter} (Başlıksız)"
            header_text = ""
            
        columns_info.append({
            "letter": letter,
            "header": header_text,
            "display": display_text,
            "index": col_idx,
        })
        
    return columns_info, df_data


def import_excel_data(
    file_path: str,
    fields_dict: dict[str, dict[str, Any]],
    field_mapping: dict[str, str],  # {db_field: excel_column_letter} (Örn: {"sku": "A"})
    db: Session,
    entity_type: str,  # "customer" veya "product"
    conflict_exist: str,  # "update" veya "skip"
    conflict_not_exist: str,  # "insert" veya "skip"
    site_id: int | None = None,
    sheet_index: int = 1,
    has_headers: bool = True,
) -> tuple[int, int, int]:  # (added_count, updated_count, skipped_count)
    """Excel/ODS dosyasını okuyarak veritabanına aktarımı gerçekleştirir ve ChangeLog üretir."""
    columns_info, df = read_excel_or_ods(file_path, sheet_index, has_headers)
    
    added_count = 0
    updated_count = 0
    skipped_count = 0
    
    df = df.fillna("")
    
    for _, row in df.iterrows():
        # Satır verisini db_field -> value sözlüğüne dönüştür
        row_data = {}
        for db_field, col_letter in field_mapping.items():
            col_idx = letter_to_col_idx(col_letter)
            if col_idx in row:
                row_data[db_field] = row[col_idx]
            else:
                row_data[db_field] = None

        # Boş satır kontrolü: Eğer tüm maplenen alanlar boş ise atla
        if all(str(val).strip() == "" for val in row_data.values() if val is not None):
            continue

        if entity_type == "customer":
            # Cari için çakışma tespiti: Vergi No veya Cari Kodu baz alınır
            tax_num = str(row_data.get("tax_number", "")).strip() if row_data.get("tax_number") else ""
            cust_code = str(row_data.get("customer_code", "")).strip() if row_data.get("customer_code") else ""
            
            existing_cust = None
            if tax_num or cust_code:
                query = db.query(Customer).filter(Customer.is_deleted == False)
                filters = []
                if tax_num:
                    filters.append(Customer.tax_number == tax_num)
                if cust_code:
                    filters.append(Customer.customer_code == cust_code)
                
                from sqlalchemy import or_
                existing_cust = query.filter(or_(*filters) if len(filters) > 1 else filters[0]).first()
                
            if existing_cust:
                if conflict_exist == "update":
                    # Güncelleme
                    for k, v in row_data.items():
                        if v is not None:
                            setattr(existing_cust, k, v)
                    db.add(existing_cust)
                    db.flush()
                    
                    # Eğer site_id tanımlıysa (Uzak sunucuya senkronizasyon seçilmişse) ChangeLog oluştur
                    if site_id is not None:
                        changelog = ChangeLog(
                            entity_type="customer",
                            entity_id=existing_cust.id,
                            action="update",
                            status="PENDING_PUSH",
                            retry_count=0,
                        )
                        db.add(changelog)
                    updated_count += 1
                else:
                    skipped_count += 1
            else:
                if conflict_not_exist == "insert":
                    new_cust = Customer(
                        marketplace="local",
                        fullname=row_data.get("fullname", "İsimsiz Cari"),
                        email=row_data.get("email"),
                        phone=row_data.get("phone"),
                        address=row_data.get("address"),
                        tax_office=row_data.get("tax_office"),
                        tax_number=row_data.get("tax_number"),
                        customer_code=row_data.get("customer_code"),
                        group_name=row_data.get("group_name"),
                        sub_group_1=row_data.get("sub_group_1"),
                        sub_group_2=row_data.get("sub_group_2"),
                        special_code_1=row_data.get("special_code_1"),
                        special_code_2=row_data.get("special_code_2"),
                        special_code_3=row_data.get("special_code_3"),
                        status=1,
                    )
                    db.add(new_cust)
                    db.flush()
                    
                    # Eğer site_id tanımlıysa (Uzak sunucuya senkronizasyon seçilmişse) ChangeLog oluştur
                    if site_id is not None:
                        changelog = ChangeLog(
                            entity_type="customer",
                            entity_id=new_cust.id,
                            action="create",
                            status="PENDING_PUSH",
                            retry_count=0,
                        )
                        db.add(changelog)
                    added_count += 1
                else:
                    skipped_count += 1
                    
        elif entity_type == "product":
            # Ürün için çakışma tespiti: SKU (Ürün Kodu) baz alınır
            sku = str(row_data.get("sku", "")).strip()
            if not sku:
                skipped_count += 1
                continue
                
            existing_prod = db.query(Product).filter(Product.sku == sku, Product.is_deleted == False).first()
            
            # Fiyat ve stok sayısal tiplerine güvenli dönüşüm
            try:
                price_val = float(row_data.get("price", 0.0)) if row_data.get("price") else 0.0
            except (ValueError, TypeError):
                price_val = 0.0
                
            try:
                base_price_val = float(row_data.get("base_price", 0.0)) if row_data.get("base_price") else 0.0
            except (ValueError, TypeError):
                base_price_val = 0.0
                
            try:
                stock_val = int(row_data.get("stock", 0)) if row_data.get("stock") else 0
            except (ValueError, TypeError):
                stock_val = 0
                
            try:
                vat_val = float(row_data.get("vat_rate", 0.0)) if row_data.get("vat_rate") else 0.0
            except (ValueError, TypeError):
                vat_val = 0.0

            if existing_prod:
                if conflict_exist == "update":
                    existing_prod.name = row_data.get("name") or existing_prod.name
                    existing_prod.description = row_data.get("description") or existing_prod.description
                    existing_prod.price = price_val
                    existing_prod.base_price = base_price_val
                    existing_prod.stock = stock_val
                    existing_prod.barcode = row_data.get("barcode") or existing_prod.barcode
                    existing_prod.vat_rate = vat_val
                    if row_data.get("status") is not None:
                        try:
                            existing_prod.status = int(row_data.get("status"))
                        except ValueError:
                            pass
                    if row_data.get("status_buy") is not None:
                        try:
                            existing_prod.status_buy = int(row_data.get("status_buy"))
                        except ValueError:
                            pass
                    # site_id varsa güncelle, yoksa local kalsın (None)
                    existing_prod.site_id = site_id
                        
                    db.add(existing_prod)
                    db.flush()
                    
                    # Eğer site_id tanımlıysa ChangeLog oluştur
                    if site_id is not None:
                        changelog = ChangeLog(
                            entity_type="product",
                            entity_id=existing_prod.id,
                            action="update",
                            status="PENDING_PUSH",
                            retry_count=0,
                        )
                        db.add(changelog)
                    updated_count += 1
                else:
                    skipped_count += 1
            else:
                if conflict_not_exist == "insert":
                    new_prod = Product(
                        site_id=site_id,
                        sku=sku,
                        name=row_data.get("name") or "İsimsiz Ürün",
                        description=row_data.get("description"),
                        base_price=base_price_val,
                        price=price_val,
                        stock=stock_val,
                        barcode=row_data.get("barcode"),
                        vat_rate=vat_val,
                        status=int(row_data.get("status")) if row_data.get("status") is not None else 1,
                        status_buy=int(row_data.get("status_buy")) if row_data.get("status_buy") is not None else 1,
                    )
                    db.add(new_prod)
                    db.flush()
                    
                    # Eğer site_id tanımlıysa ChangeLog oluştur
                    if site_id is not None:
                        changelog = ChangeLog(
                            entity_type="product",
                            entity_id=new_prod.id,
                            action="create",
                            status="PENDING_PUSH",
                            retry_count=0,
                        )
                        db.add(changelog)
                    added_count += 1
                else:
                    skipped_count += 1
                    
    db.commit()
    return added_count, updated_count, skipped_count
