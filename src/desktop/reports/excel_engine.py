"""
ToyaUI — Excel Raporlama Motoru
openpyxl ile Excel üretimi.

Kullanım:
    engine = ExcelEngine()
    wb = engine.render_teklif(data)
    engine.save(wb, "/path/to/output.xlsx")
    engine.open(wb)
"""

import logging
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side
    )
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False
    logger.warning("openpyxl bulunamadı: pip install openpyxl")


class ExcelEngine:
    """openpyxl tabanlı Excel raporlama motoru."""

    # Renkler
    COLOR_PRIMARY   = "1E3A8A"  # Koyu mavi
    COLOR_HEADER    = "2563EB"  # Açık mavi
    COLOR_ALT_ROW   = "F0F4F8"  # Açık gri
    COLOR_TOTAL     = "EFF6FF"  # Çok açık mavi
    COLOR_GRAND     = "1E3A8A"  # Genel toplam

    def _check(self):
        if not OPENPYXL_OK:
            raise RuntimeError("openpyxl yüklü değil: pip install openpyxl")

    @staticmethod
    def _fmt_currency(value) -> str:
        try:
            v = float(value or 0)
            return f"{v:,.2f} ₺".replace(",", "X").replace(".", ",").replace("X", ".")
        except Exception:
            return "0,00 ₺"

    def _header_font(self, bold=True, size=10, color="FFFFFF"):
        return Font(name="Calibri", bold=bold, size=size, color=color)

    def _data_font(self, bold=False, size=10, color="0F172A"):
        return Font(name="Calibri", bold=bold, size=size, color=color)

    def _fill(self, color):
        return PatternFill("solid", fgColor=color)

    def _border(self):
        thin = Side(style="thin", color="CBD5E1")
        return Border(left=thin, right=thin, top=thin, bottom=thin)

    def _center(self, wrap=False):
        return Alignment(
            horizontal="center", vertical="center", wrap_text=wrap
        )

    def _right(self):
        return Alignment(horizontal="right", vertical="center")

    def _left(self, wrap=False):
        return Alignment(horizontal="left", vertical="center", wrap_text=wrap)

    # ─────────────────────────────────────────────
    # TEKLİF EXCEL
    # ─────────────────────────────────────────────

    def render_teklif(self, data: dict[str, Any]) -> "Workbook":
        """Teklif verilerinden Excel workbook oluştur."""
        self._check()
        wb = Workbook()
        ws = wb.active
        ws.title = "Teklif"

        firma   = data.get("firma", {})
        musteri = data.get("musteri", {})
        belge   = data.get("belge", {})
        kalemler = data.get("kalemler", [])
        toplamlar = data.get("toplamlar", {})
        notlar  = data.get("notlar", "")

        row = 1

        # ── BAŞLIK ──
        ws.merge_cells(f"A{row}:I{row}")
        cell = ws[f"A{row}"]
        cell.value = firma.get("adi", "TOYA ERP")
        cell.font = Font(name="Calibri", bold=True, size=16, color=self.COLOR_PRIMARY)
        cell.alignment = self._left()
        ws.row_dimensions[row].height = 30
        row += 1

        ws.merge_cells(f"A{row}:I{row}")
        cell = ws[f"A{row}"]
        teklif_no = belge.get("teklif_no", "")
        cell.value = f"TEKLİF — {teklif_no}"
        cell.font = Font(name="Calibri", bold=True, size=13, color="2563EB")
        cell.alignment = self._left()
        ws.row_dimensions[row].height = 22
        row += 1

        row += 1  # boş satır

        # ── FİRMA + MÜŞTERİ BİLGİLERİ ──
        ws[f"A{row}"] = "Kimden:"
        ws[f"A{row}"].font = self._data_font(bold=True, color=self.COLOR_PRIMARY)
        ws[f"E{row}"] = "Kime:"
        ws[f"E{row}"].font = self._data_font(bold=True, color=self.COLOR_PRIMARY)
        row += 1

        firma_satir = [
            firma.get("adi", ""),
            firma.get("adres", ""),
            firma.get("tel", ""),
            firma.get("email", ""),
        ]
        musteri_satir = [
            musteri.get("adi", ""),
            musteri.get("adres", ""),
            musteri.get("vergi_daire", "") + " / " + musteri.get("vergi_no", ""),
        ]

        max_len = max(len(firma_satir), len(musteri_satir))
        for i in range(max_len):
            if i < len(firma_satir) and firma_satir[i]:
                ws[f"A{row}"] = firma_satir[i]
                ws[f"A{row}"].font = self._data_font()
            if i < len(musteri_satir) and musteri_satir[i]:
                ws[f"E{row}"] = musteri_satir[i]
                ws[f"E{row}"].font = self._data_font()
            row += 1

        row += 1  # boş satır

        # ── BELGE BİLGİLERİ ──
        belge_bilgiler = [
            ("Teklif No:",  belge.get("teklif_no", "")),
            ("Tarih:",      belge.get("tarih", "")),
            ("Vade:",       belge.get("vade", "")),
            ("Para Birimi:",belge.get("para_birimi", "TRY")),
        ]
        for label, value in belge_bilgiler:
            ws[f"A{row}"] = label
            ws[f"A{row}"].font = self._data_font(bold=True)
            ws[f"B{row}"] = value
            ws[f"B{row}"].font = self._data_font()
            row += 1

        row += 1  # boş satır

        # ── KALEMLER BAŞLIK ──
        headers = [
            "#", "Stok Kodu", "Açıklama / Ürün Adı", "Miktar",
            "Birim", "Birim Fiyat", "İsk %", "KDV %", "Tutar"
        ]
        col_widths = [5, 14, 40, 10, 10, 14, 8, 8, 16]

        for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
            col_letter = get_column_letter(ci)
            ws.column_dimensions[col_letter].width = w
            cell = ws.cell(row=row, column=ci, value=h)
            cell.font = self._header_font()
            cell.fill = self._fill(self.COLOR_HEADER)
            cell.alignment = self._center()
            cell.border = self._border()
        ws.row_dimensions[row].height = 22
        row += 1

        # ── KALEM SATIRLARI ──
        for idx, kalem in enumerate(kalemler, 1):
            bg = self.COLOR_ALT_ROW if idx % 2 == 0 else "FFFFFF"
            satir_data = [
                idx,
                kalem.get("kod", ""),
                kalem.get("aciklama", ""),
                kalem.get("miktar", 0),
                kalem.get("birim", "Adet"),
                kalem.get("birim_fiyat", 0),
                kalem.get("iskonto", 0),
                kalem.get("kdv", 20),
                kalem.get("tutar", 0),
            ]
            alignments = [
                self._center(), self._left(), self._left(wrap=True),
                self._center(), self._center(),
                self._right(), self._center(), self._center(), self._right()
            ]
            for ci, (val, aln) in enumerate(zip(satir_data, alignments), 1):
                cell = ws.cell(row=row, column=ci, value=val)
                cell.font = self._data_font()
                cell.fill = self._fill(bg)
                cell.alignment = aln
                cell.border = self._border()
                # Para formatı
                if ci in (6, 9):
                    cell.number_format = '#,##0.00 "₺"'
                if ci in (7, 8):
                    cell.number_format = '0.0"%"'
            ws.row_dimensions[row].height = 18
            row += 1

        row += 1  # boş satır

        # ── TOPLAMLAR ──
        toplam_satirlar = [
            ("Ara Toplam:",    toplamlar.get("ara_toplam", 0)),
            ("İskonto (-)",    toplamlar.get("iskonto", 0)),
            ("Masraflar (+)",  toplamlar.get("masraflar", 0)),
            ("KDV Toplam:",    toplamlar.get("kdv_toplam", 0)),
        ]
        for label, value in toplam_satirlar:
            ws.merge_cells(f"G{row}:H{row}")
            cell_l = ws[f"G{row}"]
            cell_l.value = label
            cell_l.font = self._data_font(bold=True)
            cell_l.alignment = self._right()
            cell_l.fill = self._fill(self.COLOR_TOTAL)
            cell_l.border = self._border()

            cell_v = ws[f"I{row}"]
            cell_v.value = value
            cell_v.font = self._data_font()
            cell_v.alignment = self._right()
            cell_v.fill = self._fill(self.COLOR_TOTAL)
            cell_v.border = self._border()
            cell_v.number_format = '#,##0.00 "₺"'
            row += 1

        # Genel Toplam
        ws.merge_cells(f"G{row}:H{row}")
        cell_gt_l = ws[f"G{row}"]
        cell_gt_l.value = "GENEL TOPLAM:"
        cell_gt_l.font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
        cell_gt_l.fill = self._fill(self.COLOR_GRAND)
        cell_gt_l.alignment = self._right()
        cell_gt_l.border = self._border()

        cell_gt_v = ws[f"I{row}"]
        cell_gt_v.value = toplamlar.get("genel_toplam", 0)
        cell_gt_v.font = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
        cell_gt_v.fill = self._fill(self.COLOR_GRAND)
        cell_gt_v.alignment = self._right()
        cell_gt_v.border = self._border()
        cell_gt_v.number_format = '#,##0.00 "₺"'
        ws.row_dimensions[row].height = 22
        row += 2

        # ── NOTLAR ──
        if notlar:
            ws[f"A{row}"] = "Notlar & Şartlar:"
            ws[f"A{row}"].font = self._data_font(bold=True, color=self.COLOR_PRIMARY)
            row += 1
            ws.merge_cells(f"A{row}:I{row+2}")
            cell_not = ws[f"A{row}"]
            cell_not.value = notlar
            cell_not.font = self._data_font(size=9, color="475569")
            cell_not.alignment = self._left(wrap=True)
            ws.row_dimensions[row].height = 50

        # Yazdırma ayarları
        ws.print_title_rows = "1:1"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_margins.left   = 0.5
        ws.page_margins.right  = 0.5
        ws.page_margins.top    = 0.75
        ws.page_margins.bottom = 0.75

        return wb

    # ─────────────────────────────────────────────
    # KAYDET / AÇ
    # ─────────────────────────────────────────────

    def save(self, wb: "Workbook", output_path: str) -> str:
        """Workbook'u dosyaya kaydet."""
        self._check()
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(path))
        logger.info(f"Excel kaydedildi: {path}")
        return str(path)

    def open(self, wb: "Workbook") -> str:
        """Excel'i geçici dosyaya yaz ve sistem ile aç."""
        self._check()
        tmp = tempfile.NamedTemporaryFile(
            suffix=".xlsx", delete=False, prefix="toya_"
        )
        tmp.close()
        wb.save(tmp.name)

        if sys.platform == "win32":
            os.startfile(tmp.name)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", tmp.name])
        else:
            subprocess.Popen(["xdg-open", tmp.name])

        return tmp.name
