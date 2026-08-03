"""Excel Exporter Service for Quotations and Orders."""

import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

logger = logging.getLogger(__name__)


class ExcelExporter:
    """Exports Quotations and Orders into styled Excel (.xlsx) files."""

    @staticmethod
    def export_quotation_to_excel(file_path: str, quotation: Any) -> bool:
        """Exports a Quotation or Order instance into a formatted Excel sheet."""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Teklif Formu"

            # Colors & Fonts
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            title_font = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
            bold_font = Font(name="Segoe UI", size=10, bold=True)
            regular_font = Font(name="Segoe UI", size=10)

            thin_border = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1"),
            )

            # Title
            ws.merge_cells("A1:G1")
            ws["A1"] = f"TEKLİF / SİPARİŞ FORMU - {quotation.quotation_number}"
            ws["A1"].font = title_font
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 35

            # Customer & Details Info
            cust_name = quotation.customer.fullname if quotation.customer else (quotation.customer_name_free or "Müşteri Belirtilmedi")
            tax_off = quotation.customer.tax_office if quotation.customer else (quotation.tax_office_free or "-")
            tax_num = quotation.customer.tax_number if quotation.customer else (quotation.tax_number_free or "-")

            ws["A3"] = "Müşteri / Cari Unvanı:"
            ws["A3"].font = bold_font
            ws["B3"] = cust_name
            ws["B3"].font = regular_font

            ws["A4"] = "Vergi Dairesi / No:"
            ws["A4"].font = bold_font
            ws["B4"] = f"{tax_off} / {tax_num}"
            ws["B4"].font = regular_font

            ws["E3"] = "Tarih:"
            ws["E3"].font = bold_font
            ws["F3"] = str(quotation.issue_date.strftime("%d.%m.%Y") if hasattr(quotation.issue_date, "strftime") else quotation.issue_date)
            ws["F3"].font = regular_font

            ws["E4"] = "Para Birimi:"
            ws["E4"].font = bold_font
            ws["F4"] = getattr(quotation, "currency", "TRY")
            ws["F4"].font = regular_font

            # Line items headers
            headers = ["No", "Ürün / Hizmet Açıklaması", "Birim", "Miktar", "Birim Fiyat", "KDV %", "Toplam Fiyat"]
            start_row = 6

            for col_num, h_text in enumerate(headers, 1):
                cell = ws.cell(row=start_row, column=col_num, value=h_text)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.row_dimensions[start_row].height = 25

            # Items
            current_row = start_row + 1
            for idx, item in enumerate(quotation.items, 1):
                ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center")
                ws.cell(row=current_row, column=2, value=item.product_name_free or (item.product.name if item.product else ""))
                ws.cell(row=current_row, column=3, value=item.unit or "Adet").alignment = Alignment(horizontal="center")
                ws.cell(row=current_row, column=4, value=item.quantity).alignment = Alignment(horizontal="right")
                ws.cell(row=current_row, column=5, value=item.unit_price).number_format = "#,##0.00"
                ws.cell(row=current_row, column=6, value=item.vat_rate).alignment = Alignment(horizontal="center")
                ws.cell(row=current_row, column=7, value=item.total_price).number_format = "#,##0.00"

                for col_num in range(1, 8):
                    c = ws.cell(row=current_row, column=col_num)
                    c.font = regular_font
                    c.border = thin_border

                current_row += 1

            # Totals
            current_row += 1
            totals_data = [
                ("Ara Toplam:", quotation.subtotal),
                ("Toplam İskonto:", quotation.discount_total),
                ("Toplam KDV:", quotation.vat_total),
                ("Genel Toplam:", quotation.grand_total),
            ]

            for label, val in totals_data:
                ws.cell(row=current_row, column=6, value=label).font = bold_font
                ws.cell(row=current_row, column=6).alignment = Alignment(horizontal="right")
                cell_val = ws.cell(row=current_row, column=7, value=val)
                cell_val.font = bold_font
                cell_val.number_format = "#,##0.00 ₺"
                cell_val.border = thin_border
                current_row += 1

            # Column Widths
            col_widths = [8, 35, 12, 12, 16, 12, 18]
            for col_idx, width in enumerate(col_widths, 1):
                col_letter = chr(64 + col_idx)
                ws.column_dimensions[col_letter].width = width

            wb.save(file_path)
            return True
        except Exception as e:
            logger.error(f"Error exporting quotation to Excel: {e}")
            return False
