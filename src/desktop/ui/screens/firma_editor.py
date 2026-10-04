"""
TOYA ERP - Firma (Şirket) Tanım Editörü

`DefinitionEditorScreen` (tpl.card.001 v2 — çok kolonlu bölümlü düzen) ile
Cari/Stok Kartı ile aynı kalıpta kurulan tam ekran firma kartı: numaralı
bölümler, sol panelde "BÖLÜMLER" nav listesi gizli (sekmeler zaten aynı
gezinmeyi sağlıyor) — yalnızca "İŞLEMLER" grubu altında her bölümün kendi
hızlı ekleme butonu (Dönem / Yetkili / Ortak / Banka Hesabı), sağda "Ara"
(arama) ve "Yazdır & Aktarım" (Excel/PDF/İçe Aktar/Rapor) panelleri —
`FilterWidget` + `ExportWidget` bileşenlerinden kurulur, altta VAZGEÇ /
KAYDET (+ Kaydet ve Yeni).

Dönemler, Yetkililer, Ortaklar ve Banka Hesapları artık çoklu kayıt destekleyen
"subtable" alanlarıdır (`DefinitionEditorScreen._SubTableField` — belge kalem/
masraf gridleriyle aynı standart `FilterableTableView` çekirdeğini kullanır);
modelde JSON metin (`periods_json`, `authorized_persons_json`,
`partners_json`, `bank_accounts_json`) olarak saklanır.
"""

from __future__ import annotations

import logging

from PyQt6.QtCore import Qt
from PyQt6.QtGui import QStandardItem
from PyQt6.QtWidgets import (
    QDialog,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QInputDialog,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QVBoxLayout,
    QWidget,
)

from src.desktop.services.company_service import CompanyService
from src.desktop.services.currency_catalog import CURRENCY_CODES
from src.desktop.ui.components.dialog_style import apply_light_dialog_style
from src.desktop.ui.components.entry_table_view import standart_grid
from src.desktop.ui.screens.definition_editor_screen import DefinitionEditorScreen
from src.desktop.ui.widgets.export_widget import ExportWidget
from src.desktop.ui.widgets.filter_widget import FilterWidget

logger = logging.getLogger(__name__)


class _KayitEkleDialog(QDialog):
    """Şube/Depo gibi küçük tanım kayıtları için genel amaçlı, tek seviyeli metin alanlı ekleme popup'ı."""

    def __init__(self, baslik: str, alanlar: list[tuple[str, str]], parent=None,
                 deger: dict | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle(baslik)
        self.setMinimumWidth(340)
        apply_light_dialog_style(self)
        f = QFormLayout()
        self._girdiler: dict[str, QLineEdit] = {}
        for key, label in alanlar:
            le = QLineEdit()
            if deger:
                le.setText(str(deger.get(key, "") or ""))
            f.addRow(label, le)
            self._girdiler[key] = le
        lyt = QVBoxLayout(self)
        lyt.addLayout(f)
        bar = QHBoxLayout()
        bar.addStretch()
        b_kaydet, b_iptal = QPushButton("💾 Kaydet"), QPushButton("✖ Vazgeç")
        b_kaydet.setDefault(True)
        bar.addWidget(b_kaydet)
        bar.addWidget(b_iptal)
        lyt.addLayout(bar)
        b_kaydet.clicked.connect(self.accept)
        b_iptal.clicked.connect(self.reject)

    def deger(self) -> dict:
        return {k: le.text().strip() for k, le in self._girdiler.items()}

COMPANY_TYPES = ["Şahıs", "Tüzel", "Limited Şirket", "Anonim Şirket", "Kooperatif", "Diğer"]
#: Henüz hiç döviz tanımlanmamışsa (Genel Ayarlar > Döviz & Kur boşsa) "Döviz
#: Birimi" combosunun düşeceği yedek liste — ISO 4217 kataloğunun tamamı (bkz.
#: `currency_catalog.py`), en sık kullanılanlar (TRY/USD/EUR/GBP) başta.
CURRENCIES = CURRENCY_CODES

# Çoklu kayıt alanları: {form_key: (Excel sayfa adı, [{"key","label"}...])}
_GRID_GROUPS: dict[str, tuple[str, list[dict]]] = {
    "periods": ("Dönemler", [
        {"key": "ad", "label": "Dönem Adı"},
        {"key": "baslangic", "label": "Başlangıç Tarihi", "type": "date"},
        {"key": "bitis", "label": "Bitiş Tarihi", "type": "date"},
        {"key": "aktif", "label": "Aktif (Evet/Hayır)", "type": "combo", "options": ["Evet", "Hayır"]},
    ]),
    "authorized_persons": ("Yetkililer", [
        {"key": "ad_soyad", "label": "Ad Soyad"},
        {"key": "unvan", "label": "Ünvanı"},
        {"key": "telefon", "label": "Telefon"},
        {"key": "e_posta", "label": "E-Posta"},
    ]),
    "partners": ("Ortaklar", [
        {"key": "ad_soyad", "label": "Ad Soyad"},
        {"key": "tc_vergi_no", "label": "TC / Vergi No"},
        {"key": "pay_orani", "label": "Pay Oranı (%)"},
        {"key": "gorev", "label": "Görevi"},
    ]),
    "bank_accounts": ("Banka Hesapları", [
        {"key": "banka", "label": "Banka"},
        {"key": "sube", "label": "Şube"},
        {"key": "hesap_no", "label": "Hesap No"},
        {"key": "iban", "label": "IBAN"},
        {"key": "para_birimi", "label": "Para Birimi"},
    ]),
}


def _build_firma_tabs(db_session) -> list[dict]:
    """Para Birimi alanı Döviz Tanımları'nda (Genel Ayarlar > Şirket
    Tanımları) kayıtlı aktif dövizleri dinamik olarak listeler."""

    def _load_active_currencies() -> list[str]:
        if not db_session:
            return CURRENCIES
        try:
            from sqlalchemy import select

            from src.core.models import Currency
            codes = db_session.scalars(
                select(Currency.code).where(
                    Currency.is_deleted == False, Currency.is_active == True,  # noqa: E712
                ).order_by(Currency.code),
            ).all()
            return list(codes) or CURRENCIES
        except Exception:  # noqa: BLE001
            return CURRENCIES

    def _load_countries() -> list[str]:
        try:
            from src.desktop.services.master_definitions_store import (
                MasterDefinitionsStore,
            )
            countries = MasterDefinitionsStore().get_countries(active_only=True)
            return [c["name"] for c in countries] or ["TÜRKİYE"]
        except Exception:  # noqa: BLE001
            return ["TÜRKİYE"]

    def _load_cities() -> list[str]:
        try:
            from src.desktop.services.master_definitions_store import (
                MasterDefinitionsStore,
            )
            cities = MasterDefinitionsStore().get_cities(country_code="TR", active_only=True)
            return [c["name"] for c in cities]
        except Exception:  # noqa: BLE001
            return []

    return [
        {
            "title": "Kimlik",
            "bands": [
                {"sections": [
                    {"title": "1. Firma Kimliği", "columns": 2, "stretch": 6, "fields": [
                        {"key": "code", "label": "Firma Kodu", "type": "text",
                         "required": True, "lock_on_edit": True, "placeholder": "örn. 000"},
                        {"key": "company_type", "label": "Şirket Tipi", "type": "combo",
                         "options": COMPANY_TYPES},
                        {"key": "short_name", "label": "Kısa Ad", "type": "text",
                         "required": True},
                        {"key": "default_currency", "label": "Çalışma Para Birimi", "type": "combo",
                         "options_provider": _load_active_currencies},
                        {"key": "title", "label": "Resmi / Ticari Ünvan", "type": "text",
                         "span": 2},
                        {"key": "notes", "label": "Not", "type": "multiline", "span": 2,
                         "height": 50},
                    ]},
                    {"title": "2. Vergi & Sicil Bilgileri", "columns": 2, "stretch": 6,
                     "fields": [
                        {"key": "tax_office", "label": "Vergi Dairesi", "type": "text"},
                        {"key": "tax_office_code", "label": "V.D. Kodu", "type": "text"},
                        {"key": "tax_number", "label": "Vergi No / TCKN", "type": "text"},
                        {"key": "mersis_no", "label": "MERSİS No", "type": "text"},
                        {"key": "trade_registry_no", "label": "Ticaret Sicil No", "type": "text"},
                        {"key": "nace_code", "label": "NACE / Faaliyet Kodu", "type": "text"},
                        {"key": "sgk_no", "label": "SGK İşyeri Sicil No", "type": "text"},
                        {"key": "founded_at", "label": "Kuruluş Tarihi", "type": "date"},
                    ]},
                ]},
                {"sections": [
                    {"title": "3. Adres ve İletişim", "columns": 2, "stretch": 6, "fields": [
                        {"key": "address", "label": "Açık Adres", "type": "multiline", "span": 2, "height": 50},
                        {"key": "country", "label": "Ülke", "type": "combo",
                         "options_provider": _load_countries},
                        {"key": "city", "label": "İl (Şehir)", "type": "combo",
                         "options_provider": _load_cities},
                        {"key": "district", "label": "İlçe", "type": "text"},
                        {"key": "postal_code", "label": "Posta Kodu", "type": "text"},
                        {"key": "phone1", "label": "Telefon 1", "type": "text"},
                        {"key": "phone2", "label": "Telefon 2", "type": "text"},
                        {"key": "fax", "label": "Faks", "type": "text"},
                        {"key": "email", "label": "Genel E-Posta", "type": "text"},
                        {"key": "accounting_email", "label": "Muhasebe / e-Fatura E-Posta", "type": "text"},
                        {"key": "website", "label": "Web Sitesi", "type": "text"},
                        {"key": "kep_address", "label": "KEP Adresi", "type": "text"},
                    ]},
                ]},
                {"sections": [
                    {"title": "4. Dönemler", "columns": 1, "fields": [
                        {"key": "periods", "label": "Dönemler", "type": "subtable",
                         "columns": _GRID_GROUPS["periods"][1],
                         "delete_confirm": "Seçili dönem(ler) ve onlara ait tüm bilgiler silinecek "
                                           "(kayıt veritabanında gizlenir, kalıcı silinmez). Devam edilsin mi?"},
                    ]},
                    {"title": "5. Logo & Kaşe / İmza", "columns": 2, "fields": [
                        {"key": "logo_base64", "label": "Logo (baskı)", "type": "image"},
                        {"key": "stamp_base64", "label": "Kaşe / İmza", "type": "image"},
                    ]},
                ]},
            ],
        },
        {
            "title": "Yetkililer, Ortaklar, e-Belge & Banka",
            "bands": [
                {"sections": [
                    {"title": "1. Yetkililer", "columns": 1, "fields": [
                        {"key": "authorized_persons", "label": "Yetkililer", "type": "subtable",
                         "columns": _GRID_GROUPS["authorized_persons"][1]},
                    ]},
                ]},
                {"sections": [
                    {"title": "2. Ortaklar", "columns": 1, "fields": [
                        {"key": "partners", "label": "Ortaklar", "type": "subtable",
                         "columns": _GRID_GROUPS["partners"][1]},
                    ]},
                ]},
                {"sections": [
                    {"title": "3. e-Belge", "columns": 2, "fields": [
                        {"key": "e_invoice_enabled", "label": "e-Fatura Mükellefi", "type": "bool",
                         "text": "Evet"},
                        {"key": "e_archive_enabled", "label": "e-Arşiv Mükellefi", "type": "bool",
                         "text": "Evet"},
                        {"key": "e_dispatch_enabled", "label": "e-İrsaliye Mükellefi", "type": "bool",
                         "text": "Evet"},
                        {"key": "gib_alias", "label": "GİB Etiketi / Alias", "type": "text"},
                        {"key": "integrator", "label": "Entegratör", "type": "text", "span": 2},
                    ]},
                ]},
                {"sections": [
                    {"title": "4. Banka Hesapları", "columns": 1, "fields": [
                        {"key": "bank_accounts", "label": "Banka Hesapları", "type": "subtable",
                         "columns": _GRID_GROUPS["bank_accounts"][1]},
                    ]},
                ]},
            ],
        },
    ]


# Geriye dönük uyumluluk — dinamik döviz listesi olmadan sekme şeması.
FIRMA_TABS: list[dict] = _build_firma_tabs(None)


class FirmaEditorDialog(QDialog):
    """Firma editörünü maksimize bir dialog içinde açar."""

    def __init__(self, db_session=None, company_id: int | None = None, parent=None):
        super().__init__(parent)
        self.db = db_session
        self.company_id = company_id
        self.service = CompanyService(db_session=db_session)
        self.saved_company_id: int | None = None

        self.has_saved: bool = False

        self.setWindowTitle("Firma Detayı")
        lyt = QVBoxLayout(self)
        lyt.setContentsMargins(0, 0, 0, 0)

        self.side_panel = self._build_side_panel()

        self.editor = DefinitionEditorScreen(
            title="Firma Detayı",
            tabs=_build_firma_tabs(db_session),
            code_field="code",
            title_field="short_name",
            header_extra_field="company_type",
            status_field="is_active",
            nav_sidebar=True,
            show_section_nav=False,
            enable_save_and_new=True,
            summary_widget=self.side_panel,
            right_panel_width=250,
            left_panel_width=190,
            extra_actions=[
                {"label": "➕ Yeni Dönem", "callback": lambda: self._add_row("periods")},
                {"label": "➕ Yeni Yetkili", "callback": lambda: self._add_row("authorized_persons")},
                {"label": "➕ Yeni Ortak", "callback": lambda: self._add_row("partners")},
                {"label": "➕ Yeni Banka Hesabı", "callback": lambda: self._add_row("bank_accounts")},
                {"label": "💱 Döviz Tanımlarını Yönet", "callback": self._doviz_tanimlarini_yonet},
            ],
            parent=self,
        )
        lyt.addWidget(self.editor)
        self.editor.closed.connect(self._on_close_requested)
        self.editor.saved.connect(self._on_save)
        self.editor.saved_and_new.connect(self._on_save_and_new)

        self.editor.tabw.addTab(self._build_sube_depo_tab(), "Şube & Depo Tanımları")

        if company_id:
            company = self.service.get(company_id)
            if company:
                self.editor.set_data(self.service.to_dict(company), is_edit=True)
        else:
            self.editor.set_data(CompanyService.blank_company(), is_edit=False)

        self.showMaximized()
        from src.desktop.ui.components.window_fit_helper import zorla_ilk_yerlesim
        zorla_ilk_yerlesim(self)

    # ------------------------------------------------------------------
    def _build_side_panel(self) -> FilterWidget:
        """Sağ sidebar: Ara (FilterWidget) + Yazdır/Dışa-İçe Aktar (ExportWidget) —
        ikisi de paylaşılan, diğer ekranlarda da kullanılan widget'lardan gelir."""
        panel = FilterWidget(
            group_title="ARA", show_filters=["search"], initial_open=True, panel_width=220,
        )
        panel.search_box.setPlaceholderText("Dönem / Yetkili / Ortak / Banka ara...")
        panel.filter_changed.connect(lambda f: self._apply_quick_search(f.get("search", "")))
        panel.filters_cleared.connect(lambda: self._apply_quick_search(""))

        self.export_widget = ExportWidget(group_title="YAZDIR & AKTARIM", initial_open=True)
        self.export_widget.print_clicked.connect(self._on_print)
        self.export_widget.export_pdf_clicked.connect(self._on_export_pdf)
        self.export_widget.export_excel_clicked.connect(self._on_export_excel)
        self.export_widget.import_excel_clicked.connect(self._on_import_excel)
        self.export_widget.report_clicked.connect(self._on_report)
        panel.add_custom_section(self.export_widget)
        return panel

    # ------------------------------------------------------------------ Şube & Depo Tanımları (2026-09-30)
    def _bu_firma_id(self) -> int | None:
        """Bu kartın ait olduğu firma id'si — yeni (henüz kaydedilmemiş) kartta None."""
        return self.saved_company_id or self.company_id

    def _build_sube_depo_tab(self) -> QWidget:
        """Şube ve Depo Tanımları'nı ayrı bir sekme olarak gösterir — Genel Ayarlar'daki gerçek
        `SubeTanimi`/`WarehouseDefinition` tablolarından CANLI çekilir (JSON alt-tablo değil, bu
        firmaya ait gerçek kayıtlar); buradan eklenen bir şube/depo Genel Ayarlar'da da anında görünür."""
        w = QWidget()
        lyt = QVBoxLayout(w)
        lyt.setContentsMargins(12, 12, 12, 12)
        lyt.setSpacing(8)

        info = QLabel("Bu firmaya ait şube ve depo tanımları — Genel Ayarlar'daki 'Şube Tanımları' / "
                       "'Depo Tanımları' ile AYNI kayıtlardır, buradan eklenen orada da görünür.")
        info.setWordWrap(True)
        info.setStyleSheet("color:#64748b; font-size:11px;")
        lyt.addWidget(info)

        row = QHBoxLayout()
        row.setSpacing(12)

        sube_alanlar = [("code", "Şube Kodu *"), ("name", "Şube Adı *"), ("yetkili", "Yetkili")]
        depo_alanlar = [("code", "Depo Kodu *"), ("name", "Depo Adı *"), ("address", "Adres")]

        sube_col = QVBoxLayout()
        sube_bar = QHBoxLayout()
        sube_bar.addWidget(QLabel("<b>🏢 Şube Tanımları</b>"))
        sube_bar.addStretch()
        btn_sube = QPushButton("➕ Yeni Şube")
        btn_sube.clicked.connect(lambda: self._yeni_kayit_ekle("sube", "Yeni Şube", sube_alanlar))
        btn_sube_duzenle = QPushButton("✏️ Düzenle")
        btn_sube_duzenle.clicked.connect(lambda: self._kayit_duzenle("sube", "Şube Düzenle", sube_alanlar))
        btn_sube_sil = QPushButton("🗑️ Sil")
        btn_sube_sil.clicked.connect(lambda: self._kayit_sil("sube"))
        for b in (btn_sube, btn_sube_duzenle, btn_sube_sil):
            sube_bar.addWidget(b)
        sube_col.addLayout(sube_bar)
        self.grid_sube = standart_grid("firma_detay_sube", [
            ("code", "Şube Kodu"), ("name", "Şube Adı"), ("yetkili", "Yetkili"), ("durum", "Durum"),
        ])
        sube_col.addWidget(self.grid_sube)

        depo_col = QVBoxLayout()
        depo_bar = QHBoxLayout()
        depo_bar.addWidget(QLabel("<b>🏭 Depo Tanımları</b>"))
        depo_bar.addStretch()
        btn_depo = QPushButton("➕ Yeni Depo")
        btn_depo.clicked.connect(lambda: self._yeni_kayit_ekle("depo", "Yeni Depo", depo_alanlar))
        btn_depo_duzenle = QPushButton("✏️ Düzenle")
        btn_depo_duzenle.clicked.connect(lambda: self._kayit_duzenle("depo", "Depo Düzenle", depo_alanlar))
        btn_depo_sil = QPushButton("🗑️ Sil")
        btn_depo_sil.clicked.connect(lambda: self._kayit_sil("depo"))
        for b in (btn_depo, btn_depo_duzenle, btn_depo_sil):
            depo_bar.addWidget(b)
        depo_col.addLayout(depo_bar)
        self.grid_depo = standart_grid("firma_detay_depo", [
            ("code", "Depo Kodu"), ("name", "Depo Adı"), ("address", "Adres"), ("durum", "Durum"),
        ])
        depo_col.addWidget(self.grid_depo)

        row.addLayout(sube_col)
        row.addLayout(depo_col)
        lyt.addLayout(row, 1)
        self._sube_depo_yenile()
        return w

    def _sube_depo_yenile(self) -> None:
        firma_id = self._bu_firma_id()
        self.grid_sube.table_view.setRowCount(0)
        self.grid_depo.table_view.setRowCount(0)
        if not self.db or not firma_id:
            return
        from src.desktop.services.sube_service import SubeService
        for s in SubeService(self.db).liste(firma_id):
            r = self.grid_sube.table_view.rowCount()
            self.grid_sube.table_view.insertRow(r)
            for c, val in enumerate((s.code, s.name, s.yetkili or "", "Aktif" if s.is_active else "Pasif")):
                it = QStandardItem(val)
                if c == 0:
                    it.setData(s.id, Qt.ItemDataRole.UserRole)
                self.grid_sube.table_view.setItem(r, c, it)

        from sqlalchemy import select

        from src.core.models import WarehouseDefinition
        depolar = self.db.scalars(
            select(WarehouseDefinition).where(WarehouseDefinition.is_deleted == False,  # noqa: E712
                                               WarehouseDefinition.company_id == firma_id),
        ).all()
        for dp in depolar:
            r = self.grid_depo.table_view.rowCount()
            self.grid_depo.table_view.insertRow(r)
            for c, val in enumerate((dp.code, dp.name or "", dp.address or "", "Aktif" if dp.is_active else "Pasif")):
                it = QStandardItem(val)
                if c == 0:
                    it.setData(dp.id, Qt.ItemDataRole.UserRole)
                self.grid_depo.table_view.setItem(r, c, it)

    def _secili_kayit_id(self, tur: str) -> int | None:
        grid = self.grid_sube if tur == "sube" else self.grid_depo
        row = grid.table_view.currentIndex().row()
        if row < 0:
            return None
        it = grid.table_view.item(row, 0)
        return it.data(Qt.ItemDataRole.UserRole) if it else None

    def _yeni_kayit_ekle(self, tur: str, baslik: str, alanlar: list[tuple[str, str]]) -> None:
        firma_id = self._bu_firma_id()
        if not firma_id:
            QMessageBox.information(self, "Önce Kaydedin",
                                     f"{'Şube' if tur == 'sube' else 'Depo'} ekleyebilmek için önce firma kartını kaydedin.")
            return
        dlg = _KayitEkleDialog(baslik, alanlar, parent=self)
        if not dlg.exec():
            return
        v = dlg.deger()
        if not v.get("code") or not v.get("name"):
            QMessageBox.warning(self, "Eksik Bilgi", "Kod ve ad zorunludur.")
            return
        if tur == "sube":
            from src.desktop.services.sube_service import SubeService
            res = SubeService(self.db).kaydet(firma_id, v)
            if not res.success:
                QMessageBox.critical(self, "Hata", res.error or "Şube kaydedilemedi.")
                return
        else:
            from src.core.models import WarehouseDefinition
            self.db.add(WarehouseDefinition(company_id=firma_id, code=v["code"], name=v["name"],
                                             address=v.get("address") or None, is_active=True))
            self.db.commit()
        self._sube_depo_yenile()

    def _kayit_duzenle(self, tur: str, baslik: str, alanlar: list[tuple[str, str]]) -> None:
        kayit_id = self._secili_kayit_id(tur)
        if not kayit_id:
            QMessageBox.information(self, "Seçim Yok", "Önce düzenlemek istediğiniz satırı seçin.")
            return
        if tur == "sube":
            from src.desktop.services.sube_service import SubeService
            svc = SubeService(self.db)
            kayit = svc.get(kayit_id)
            if kayit is None:
                return
            onceki = {"code": kayit.code, "name": kayit.name, "yetkili": kayit.yetkili}
        else:
            from src.core.models import WarehouseDefinition
            kayit = self.db.get(WarehouseDefinition, kayit_id)
            if kayit is None:
                return
            onceki = {"code": kayit.code, "name": kayit.name, "address": kayit.address}
        dlg = _KayitEkleDialog(baslik, alanlar, parent=self, deger=onceki)
        if not dlg.exec():
            return
        v = dlg.deger()
        if not v.get("code") or not v.get("name"):
            QMessageBox.warning(self, "Eksik Bilgi", "Kod ve ad zorunludur.")
            return
        if tur == "sube":
            res = svc.kaydet(self._bu_firma_id(), v, sube_id=kayit_id)
            if not res.success:
                QMessageBox.critical(self, "Hata", res.error or "Şube güncellenemedi.")
                return
        else:
            kayit.code, kayit.name, kayit.address = v["code"], v["name"], v.get("address") or None
            self.db.commit()
        self._sube_depo_yenile()

    def _kayit_sil(self, tur: str) -> None:
        kayit_id = self._secili_kayit_id(tur)
        if not kayit_id:
            QMessageBox.information(self, "Seçim Yok", "Önce silmek istediğiniz satırı seçin.")
            return
        ad = "Şube" if tur == "sube" else "Depo"
        if QMessageBox.question(self, f"{ad} Sil", f"Seçili {ad.lower()} gizlenecek (kalıcı silinmez). "
                                 "Devam edilsin mi?") != QMessageBox.StandardButton.Yes:
            return
        if tur == "sube":
            from src.desktop.services.sube_service import SubeService
            SubeService(self.db).gizle(kayit_id)
        else:
            from src.core.models import WarehouseDefinition
            wd = self.db.get(WarehouseDefinition, kayit_id)
            if wd is not None:
                wd.is_deleted = True
                self.db.commit()
        self._sube_depo_yenile()

    def _doviz_tanimlarini_yonet(self) -> None:
        """'Çalışma Para Birimi' kısıtlı bir statik listeden değil Döviz & Kur tanımlarından geliyor;
        gerektiğinde kullanıcı o ekranın TÜM işlemlerini (ekle/sil/kur çek) buradan, ayrılmadan
        yapabilsin diye Genel Ayarlar'ın aynı paneli maksimize bir diyalogda açılır (2026-09-30).
        Kapanınca "Çalışma Para Birimi" combosu güncel aktif döviz listesiyle tazelenir."""
        from PyQt6.QtWidgets import QDialog, QVBoxLayout

        from src.desktop.ui.settings import GeneralSettingsScreen
        dlg = QDialog(self)
        dlg.setWindowTitle("Döviz & Kur Tanımları")
        lyt = QVBoxLayout(dlg)
        lyt.setContentsMargins(0, 0, 0, 0)
        scr = GeneralSettingsScreen(db_session=self.db, parent=dlg)
        scr.switch_panel("doviz_kur")
        lyt.addWidget(scr)
        dlg.showMaximized()
        dlg.exec()

        from sqlalchemy import select

        from src.core.models import Currency
        combo = self.editor._widgets.get("default_currency")  # noqa: SLF001
        if combo is None:
            return
        secili = combo.currentText()
        codes = self.db.scalars(
            select(Currency.code).where(Currency.is_deleted == False, Currency.is_active == True)  # noqa: E712
            .order_by(Currency.code),
        ).all()
        combo.blockSignals(True)
        combo.clear()
        combo.addItems(list(codes) or CURRENCIES)
        idx = combo.findText(secili)
        combo.setCurrentIndex(idx if idx >= 0 else 0)
        combo.blockSignals(False)

    def _grid_widget(self, key: str):
        return self.editor._widgets.get(key)  # noqa: SLF001

    def _add_row(self, key: str) -> None:
        if key == "periods":
            self._yeni_donem_ekle()
            return
        if key == "authorized_persons":
            self._yeni_yetkili_ekle()
            return
        if key == "partners":
            self._yeni_ortak_ekle()
            return
        if key == "bank_accounts":
            self._yeni_banka_hesabi_ekle()
            return
        w = self._grid_widget(key)
        if w:
            w._add_row({})  # noqa: SLF001

    def _yeni_donem_ekle(self) -> None:
        """'Dönemler' için takvim destekli popup form açar."""
        from src.desktop.ui.widgets.donem_ekle_dialog import DonemEkleDialog
        dlg = DonemEkleDialog(parent=self)
        if dlg.exec():
            w = self._grid_widget("periods")
            if w:
                w._add_row(dlg.deger())  # noqa: SLF001

    def _yeni_yetkili_ekle(self) -> None:
        """'Yetkililer' için popup form açar."""
        from src.desktop.ui.widgets.firma_tanim_dialoglari import YetkiliEkleDialog
        dlg = YetkiliEkleDialog(parent=self)
        if dlg.exec():
            w = self._grid_widget("authorized_persons")
            if w:
                w._add_row(dlg.deger())  # noqa: SLF001

    def _yeni_ortak_ekle(self) -> None:
        """'Ortaklar' için popup form açar."""
        from src.desktop.ui.widgets.firma_tanim_dialoglari import OrtakEkleDialog
        dlg = OrtakEkleDialog(parent=self)
        if dlg.exec():
            w = self._grid_widget("partners")
            if w:
                w._add_row(dlg.deger())  # noqa: SLF001

    def _yeni_banka_hesabi_ekle(self) -> None:
        """'Banka Hesapları' için popup form açar."""
        from src.desktop.ui.widgets.firma_tanim_dialoglari import BankaHesabiEkleDialog
        dlg = BankaHesabiEkleDialog(parent=self)
        if dlg.exec():
            w = self._grid_widget("bank_accounts")
            if w:
                w._add_row(dlg.deger())  # noqa: SLF001

    def _apply_quick_search(self, text: str) -> None:
        """Arama kutusundaki metni Dönem/Yetkili/Ortak/Banka Hesabı grid'lerinin
        tamamında filtre olarak uygular (satırların hepsini tarayıp eşleşmeyeni gizler)."""
        text = (text or "").strip().lower()
        for key in _GRID_GROUPS:
            w = self._grid_widget(key)
            if not w:
                continue
            model = w.model
            for r in range(model.rowCount()):
                cells = " ".join(
                    (model.item(r, c).text() if model.item(r, c) else "")
                    for c in range(1, model.columnCount())
                ).lower()
                w.grid.table_view.setRowHidden(r, bool(text) and text not in cells)

    # ------------------------------------------------------------------
    def _on_print(self) -> None:
        QMessageBox.information(self, "Yazdır", "Firma kartı yazdırma özelliği hazırlanıyor.")

    def _on_export_pdf(self) -> None:
        QMessageBox.information(self, "PDF", "Firma kartını PDF olarak kaydetme özelliği hazırlanıyor.")

    def _on_report(self) -> None:
        QMessageBox.information(self, "Rapor", "Detaylı firma raporu hazırlanıyor.")

    def _on_export_excel(self) -> None:
        """Dönemler, Yetkililer, Ortaklar ve Banka Hesapları listelerini tek bir
        çok sayfalı Excel dosyasına aktarır."""
        path, _ = QFileDialog.getSaveFileName(
            self, "Excel'e Aktar", "firma_listeleri.xlsx", "Excel Dosyası (*.xlsx)",
        )
        if not path:
            return
        try:
            from openpyxl import Workbook
            wb = Workbook()
            wb.remove(wb.active)
            for key, (sheet_name, columns) in _GRID_GROUPS.items():
                w = self._grid_widget(key)
                if not w:
                    continue
                ws = wb.create_sheet(sheet_name)
                ws.append([c["label"] for c in columns])
                for rec in w.value():
                    ws.append([rec.get(c["key"], "") for c in columns])
            wb.save(path)
            QMessageBox.information(self, "Dışa Aktarıldı", f"Listeler dışa aktarıldı:\n{path}")
        except Exception as exc:  # noqa: BLE001
            logger.exception("Firma listeleri Excel'e aktarılamadı")
            QMessageBox.critical(self, "Hata", f"Dışa aktarma başarısız: {exc}")

    def _on_import_excel(self) -> None:
        """Kullanıcının seçtiği listeye (Dönemler/Yetkililer/Ortaklar/Banka
        Hesapları) bir Excel dosyasından satır ekler."""
        labels = [group[0] for group in _GRID_GROUPS.values()]
        keys = list(_GRID_GROUPS.keys())
        choice, ok = QInputDialog.getItem(
            self, "İçe Aktar", "Hangi listeye aktarılsın?", labels, 0, False,
        )
        if not ok or not choice:
            return
        key = keys[labels.index(choice)]
        _, columns = _GRID_GROUPS[key]
        path, _ = QFileDialog.getOpenFileName(
            self, f"{choice} - Excel'den İçe Aktar", "", "Excel Dosyası (*.xlsx)",
        )
        if not path:
            return
        w = self._grid_widget(key)
        if not w:
            return
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            added = 0
            for row in rows:
                if not any(row):
                    continue
                rec = {c["key"]: ("" if v is None else str(v))
                       for c, v in zip(columns, row, strict=False)}
                w._add_row(rec)  # noqa: SLF001
                added += 1
            QMessageBox.information(self, "İçe Aktarıldı", f"{choice}: {added} satır eklendi.")
        except Exception as exc:  # noqa: BLE001
            logger.exception("Excel'den içe aktarma başarısız")
            QMessageBox.critical(self, "Hata", f"İçe aktarma başarısız: {exc}")

    # ------------------------------------------------------------------
    def _on_close_requested(self):
        if self.has_saved:
            self.accept()
        else:
            self.reject()

    def closeEvent(self, event):  # noqa: N802
        if self.has_saved:
            self.setResult(QDialog.DialogCode.Accepted)
        super().closeEvent(event)

    # ------------------------------------------------------------------
    def _save_payload(self, payload: dict):
        res = self.service.save(payload, company_id=self.company_id)
        if not res.success:
            QMessageBox.critical(self, "Kayıt Hatası", res.error or "Firma kaydedilemedi.")
            return None
        self.saved_company_id = res.company_id
        self.has_saved = True
        return res

    def _on_save(self, payload: dict):
        """'Kaydet' artık ekranı kapatmıyor, aynı kartta kalıyor (2026-09-30) — 'Kaydet ve Yeni' hâlâ
        formu temizliyor, ayrı bir 'Kapat' isteniyorsa pencerenin kendi kapatma (X) düğmesi/Esc kullanılır."""
        res = self._save_payload(payload)
        if res is None:
            return
        yeni_kayitti = self.company_id is None
        self.company_id = res.company_id
        if yeni_kayitti:
            w = self.editor._widgets.get("code")  # noqa: SLF001
            if w is not None and self.editor._field_specs.get("code", {}).get("lock_on_edit"):  # noqa: SLF001
                w.setEnabled(False)
        self._sube_depo_yenile()
        QMessageBox.information(self, "Kaydedildi", f"Firma kaydedildi: {res.code}")

    def _on_save_and_new(self, payload: dict):
        res = self._save_payload(payload)
        if res is None:
            return
        code = res.code
        self.company_id = None
        self.saved_company_id = None
        self.editor.set_data(CompanyService.blank_company(), is_edit=False)
        self._sube_depo_yenile()
        QMessageBox.information(
            self, "Kaydedildi", f"Firma kaydedildi: {code}. Yeni kart için form temizlendi.",
        )
